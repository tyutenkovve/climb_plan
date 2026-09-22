"""Build the private import package from the unchanged source documents."""
from __future__ import annotations
import json
from datetime import date, datetime
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'app' / 'src' / 'program-package.json'
lines = (ROOT / 'Программа.md').read_text(encoding='utf-8').splitlines()
book = openpyxl.load_workbook(ROOT / 'Bouldering_Plan_2026_2027.xlsx', data_only=True)
sections_spec = [
    ('goals', 'Цели и отправная точка', 1, 9),
    ('safety', 'Условия начала и правила остановки', 10, 18),
    ('return', 'Возвращение к нагрузке', 19, 41),
    ('stages', 'Этапы подготовки и разгрузка', 42, 60),
    ('templates', 'Основная неделя и шаблоны', 61, 76),
    ('fingers', 'Пальцы и фингерборд', 77, 101),
    ('strength', 'ОФП, ЛФК и домашняя работа', 102, 112),
    ('technique', 'Техника и тесные положения', 113, 119),
    ('moonboard', 'MoonBoard и проекты', 120, 129),
    ('endurance', 'Связки и выносливость', 130, 141),
    ('motivation', 'Мотивация и наблюдаемый прогресс', 142, 161),
    ('height', 'Практика высоты', 162, 169),
    ('recovery', 'Восстановление и изменение нагрузки', 170, 178),
    ('nutrition', 'Питание и физиология', 179, 190),
    ('progress', 'Контроль прогресса', 191, 199),
]
sections = [dict(id=i, title=t, text='\n'.join(lines[a-1:b]).strip(), source=f'Программа.md:{a}-{b}') for i,t,a,b in sections_spec]

def iso(v):
    return v.date().isoformat() if isinstance(v, datetime) else v.isoformat() if isinstance(v,date) else v

kind_by_mode = {'F':'hang','ОФП':'strength'}
def kind(mode, block):
    low = block.lower()
    if mode.startswith('F-') or 'пальцы' in low: return 'hang'
    if mode.startswith('ОФП') or mode == 'ЛФК': return 'strength'
    if 'лазание' in low or 'проект' in low or 'moonboard' in low or 'движения' in low or 'выносл' in mode.lower() or 'связки' in mode.lower(): return 'climb'
    return 'simple'

exercises=[]
for rownum in range(6,38):
    mode,block,minutes,dose,instruction = [book['Занятия'].cell(rownum,col).value for col in range(1,6)]
    if mode is None: continue
    eid=f'ex-{rownum}'
    exercises.append(dict(id=eid, mode=mode, title=block, kind=kind(mode,block), minutes=minutes, dose=dose, instruction=instruction, source=f'Занятия!A{rownum}:E{rownum}', section='fingers' if mode.startswith('F-') or block=='Пальцы — F' else 'templates'))

templates=[]
for mode,title in [('R','Возвращение R'),('A','Пальцы и проекты A'),('B','MoonBoard и техника B'),('C','Лёгкая C'),('Дома','Домашняя работа')]:
    selected=[e['id'] for e in exercises if e['mode']==mode and e['id'] not in ('ex-14','ex-19')]
    if mode=='A': selected += ['ex-26','ex-27','ex-30']
    if mode=='B': selected += ['ex-28','ex-29','ex-30']
    templates.append(dict(id=f'builtin-{mode}', name=title, type=mode, exerciseIds=selected, standard=True, version=1))
weeks=[]
for rownum in range(6,44):
    values=[book['Календарь'].cell(rownum,col).value for col in range(1,9)]
    if not isinstance(values[0],datetime): continue
    weeks.append(dict(start=iso(values[0]),end=iso(values[1]),phase=values[2],a=values[3],b=values[4],volume=values[5],note=values[6],visits=values[7]))
sources=[]
for rownum in range(6,24):
    sid,name,kind_value,note,url=[book['Источники'].cell(rownum,col).value for col in range(1,6)]
    sources.append(dict(id=sid,name=name,kind=kind_value,note=note,url=url))
start_rules=[dict(label=book['Старт'].cell(r,1).value,text=book['Старт'].cell(r,3).value,source=f'Старт!A{r}:C{r}') for r in range(11,25) if book['Старт'].cell(r,1).value and book['Старт'].cell(r,3).value]
package=dict(startRules=start_rules,id='personal-bouldering-2026',version='1.0.0',title='Боулдеринг · сентябрь 2026 — май 2027',sections=sections,exercises=exercises,templates=templates,weeks=weeks,sources=sources,glossary={
'RPE':'Субъективная тяжесть занятия от 0 (покой) до 10 (максимум).',
'Попытка':'Каждый отдельный силовой заход на движение или трассу; повторные рывки считаются отдельно.',
'Проект':'Трасса или задача, над которой работа продолжается несколько занятий.',
'Benchmark':'Отмеченная создателями MoonBoard задача в конкретной расстановке и угле.',
'Разгрузка':'Неделя с уменьшенным объёмом висов и тяжёлых попыток, без новых предельных движений.'})
assert len(weeks)==38 and len(sources)==18 and len(exercises)==32
OUT.parent.mkdir(exist_ok=True)
OUT.write_text(json.dumps(package,ensure_ascii=False,indent=2,default=iso),encoding='utf-8')
print(f'{OUT}: {len(sections)} sections, {len(exercises)} exercises, {len(weeks)} weeks, {len(sources)} sources')
